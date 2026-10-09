"""
Adapters — um por TIPO de sistema de permits, não por cidade.

Cada cidade vira só uma entrada em registry.yaml. Os adapters sabem conversar
com a plataforma (CKAN, Socrata, ArcGIS) e traduzem os campos de cada cidade
para o formato padrão do Permit Radar.

Versão 2.11 (tipo só 'residencial/comercial' sem descrição = Unspecified): Worcester (ArcGIS) e relatórios em arquivo; filtra permits de especialidade, classifica o tipo de obra
(New Construction / Addition / Renovation / Demolition) e descobre colunas sozinho.
Para ver colunas e tipos de permit no log, defina PERMIT_DEBUG=1.
"""

from __future__ import annotations

import csv
import html as html_lib
import io
import os
import re
import time
from datetime import date, datetime, timedelta, timezone
import urllib.robotparser as robotparser
from urllib.parse import urljoin, urlparse

import requests

DEBUG = bool(os.environ.get("PERMIT_DEBUG"))

HEADERS = {"User-Agent": "PermitRadar/1.0 (public permit data aggregator)"}

# Nomes comuns de colunas, em ordem de preferência. O primeiro que existir vence.
DEFAULT_CANDIDATES = {
    "permit_number": ["permit_number", "permitnumber", "permit_no", "permit_num", "permit_nbr",
                      "permit", "record_number", "record__", "record_no", "permit_id",
                      "ap_no", "record", "plannumber", "id"],
    "address": ["address", "full_address", "site_address", "street_address",
                "property_address", "project_address", "location_address", "location"],
    "permit_type": ["permit_type", "permittypedescr", "permit_type_description",
                    "permit_for", "appl_type", "record_type", "worktype", "work_type", "type"],
    "status": ["status", "permit_status", "record_status", "appl_status", "current_status"],
    "estimated_value": ["estimated_cost", "declared_valuation", "total_project_cost",
                        "total_cost_of_construction", "total_cost", "project_value",
                        "estimated_value", "est_cost", "job_cost", "project_cost",
                        "construction_cost", "cost_of_construction", "valuation",
                        "building_cost", "value", "cost"],
    "issue_date": ["issue_date", "issued_date", "date_issued", "permit_issue_date",
                   "permit_license_issued_date", "issuance_date", "issued", "appl_date", "date_submitted"],
    "description": ["description", "comments", "project_description",
                    "work_description", "brief_description", "scope_of_work", "description_of_work",
                    "detailed_description_of_work", "isd_approved_description"],
    "owner": ["owner", "owner_name", "property_owner", "owner_legal_name"],
    "applicant": ["applicant", "applicant_name"],
    "contractor": ["contractor", "contractor_name", "general_contractor", "builder",
                   "firm_name", "licensed_name", "gc_name", "licensed_contractor"],
}

ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}")
EMPTY_VALUES = {"n/a", "na", "none", "null", "-", "--", "unknown", "user", ""}


# ---------------------------------------------------------------------------
# Relevância e classificação
# O radar é para construção e reforma. Permits de especialidade (elétrico,
# hidráulico, gás, alarme...) e administrativos (certificados, emendas) ficam de fora.
# ---------------------------------------------------------------------------
INCLUDE_TYPE_RE = re.compile(
    r"(building|bldg|short form|long form|erect|new construction|foundation|alteration|"
    r"addition|demolition|remodel|renovation|construction|roof|siding|residential|commercial|"
    r"deck|accessory|adu|garage|carport|restore after fire|resi|comm\b|bldg|demo|\bsf[ad]\b|\bmfd\b|\bfnd\b)", re.I)
EXCLUDE_TYPE_RE = re.compile(
    r"\b(amendment|certificate|electrical|plumbing|gas|mechanical|sheet metal|fire alarms?|"
    r"low voltage|sprinkler|signs?|occupancy|excavation|asbestos|tents?|dumpster|food|tobacco|"
    r"insulation|stove|antenna|zoning|crowd|carnival|billboard|turbine|pool|fence|retaining|"
    r"shed|change of use|temporary|moving)\b|plumb|elect|sheet|shtmtl|sht mtl|chim\.|firealarm|"
    r"\bco\b|\bcoc\b|\bcou\b", re.I)
DESC_INCLUDE_RE = re.compile(
    r"\b(roof\w*|re-?roof\w*|siding|addition|deck|remodel\w*|renovat\w*|demoli\w*|garage|foundation|"
    r"porch|sunroom|dormer|kitchen|bath\w*|windows?|new (?:single|two|three|multi|dwelling|home|house|building|construction)|"
    r"build\w* (?:a|an|new)|construct\w*)\b", re.I)
EVENT_DESC_RE = re.compile(r"\b(tents?|beer garden|festival|one day event|1 day event)\b", re.I)

NEW_CONSTRUCTION_DESC_RE = re.compile(
    r"\bnew (building|structure|construction|single[- ]family|two[- ]family|three[- ]family|"
    r"multi[- ]?family|dwelling|home|house)\b", re.I)
WORK_ON_EXISTING_RE = re.compile(r"\b(renovat\w*|remodel\w*|alteration\w*|repair\w*|replac\w*|install\w*)\b", re.I)
DEMOLITION_RE = re.compile(r"\b(demolish\w*|demolition|demo)\b", re.I)
BUILD_WORDS_RE = re.compile(r"\b(renovat\w*|remodel\w*|replac\w*|install\w*|build|construct\w*|new|rebuild\w*)\b", re.I)
ADDITION_RE = re.compile(r"\b(addition|dormer|sunroom|bump[- ]?out)\b", re.I)


def is_relevant(permit_type, description=None) -> bool:
    """True se o permit é de construção/reforma geral (e não de especialidade ou administrativo)."""
    t = permit_type or ""
    if t:
        if EXCLUDE_TYPE_RE.search(t):
            return False
        if not INCLUDE_TYPE_RE.search(t):
            # tipo desconhecido (ex.: siglas de cada cidade): a descrição do serviço decide
            if not (description and DESC_INCLUDE_RE.search(description)):
                return False
    if description and EVENT_DESC_RE.search(description):
        return False
    return True


def classify_permit(permit_type, description=None) -> str:
    """Padroniza o tipo de obra: New Construction / Addition / Renovation / Demolition
    (ou Unspecified, quando a cidade não informa o tipo)."""
    t = (permit_type or "")
    d = (description or "")
    if re.search(r"type not listed", t, re.I):
        return "Unspecified"
    if re.fullmatch(r"\s*(sfa|sfd|mfd|sf|fnd)\.?\s*", t, re.I) and not DESC_INCLUDE_RE.search(d):
        return "Unspecified"      # sigla sem significado confirmado e sem descrição que ajude
    if not d.strip() and re.fullmatch(r"\s*(resi|comm|residential|commercial)\.?\s*", t, re.I):
        return "Unspecified"      # "residencial" ou "comercial" sozinho não diz se é obra nova, ampliação ou reforma
    if re.search(r"new construction|erect", t, re.I):
        return "New Construction"
    if NEW_CONSTRUCTION_DESC_RE.search(d) and not WORK_ON_EXISTING_RE.search(d):
        return "New Construction"
    if re.search(r"demolition", t, re.I) or (DEMOLITION_RE.search(d) and not BUILD_WORDS_RE.search(d)):
        return "Demolition"
    if re.search(r"accessory dwelling|\badu\b", t, re.I):
        return "Addition"
    if re.search(r"\baddition\b", t, re.I) and not re.search(r"alteration", t, re.I):
        return "Addition"
    if ADDITION_RE.search(d):
        return "Addition"
    return "Renovation"


_ROBOTS_CACHE: dict = {}


def robots_allows(url: str, agent: str = "PermitRadar") -> bool:
    """Respeita o robots.txt do site. Sem robots.txt acessível, segue de boa-fé."""
    p = urlparse(url)
    base = f"{p.scheme}://{p.netloc}"
    rp = _ROBOTS_CACHE.get(base)
    if rp is None:
        rp = robotparser.RobotFileParser()
        rp.set_url(base + "/robots.txt")
        try:
            rp.read()
        except Exception:  # noqa: BLE001
            rp = False
        _ROBOTS_CACHE[base] = rp
    if rp is False:
        return True
    return rp.can_fetch(agent, url)


def _http_get(url: str, params: dict | None = None):
    """GET com 3 tentativas para oscilações de rede; em erro HTTP, devolve a explicação do servidor."""
    last_exc = None
    for attempt in range(3):
        try:
            resp = requests.get(url, params=params, headers=HEADERS, timeout=60)
        except (requests.Timeout, requests.ConnectionError) as e:
            last_exc = e
            time.sleep(3 * (attempt + 1))
            continue
        if resp.status_code in (502, 503, 504) and attempt < 2:
            time.sleep(3 * (attempt + 1))
            continue
        if resp.status_code >= 400:
            raise RuntimeError(f"HTTP {resp.status_code} — resposta do servidor: {resp.text[:400]}")
        return resp
    raise RuntimeError(f"falha de rede após 3 tentativas: {last_exc}")


class BaseAdapter:
    STANDARD_FIELDS = list(DEFAULT_CANDIDATES.keys())

    def __init__(self, city: str, state: str, config: dict, field_map: dict | None, source_link: str = ""):
        self.city = city
        self.state = state
        self.config = config or {}
        self.field_map = field_map or {}
        self.source_link = source_link

    # ---- a ser implementado por cada plataforma ----
    def fetch_raw(self, days_back: int) -> list[dict]:
        raise NotImplementedError

    # ---- utilidades comuns ----
    def _pick(self, raw: dict, std_field: str):
        lower = {str(k).lower(): k for k in raw}
        configured = self.field_map.get(std_field)
        if isinstance(configured, list):
            candidates = list(configured)
        elif configured:
            candidates = [configured]
        else:
            candidates = []
        candidates += DEFAULT_CANDIDATES.get(std_field, [])
        for name in candidates:
            key = lower.get(str(name).lower())
            if key is None:
                continue
            val = raw.get(key)
            if val in (None, ""):
                continue
            # "N/A" e similares viram vazio (exceto no tipo, onde ele ajuda a descartar o registro)
            if std_field != "permit_type" and isinstance(val, str) and val.strip().lower() in EMPTY_VALUES:
                continue
            return val
        if self.config.get("prefix_match"):
            # cabeçalhos longos, ex.: "project_cost_please_enter_a_whole_number..." casam com "project_cost"
            for name in candidates:
                n = str(name).lower()
                if len(n) < 8:
                    continue
                for lk, key in lower.items():
                    if lk.startswith(n + "_") and raw.get(key) not in (None, ""):
                        val = raw[key]
                        if not (std_field != "permit_type" and isinstance(val, str)
                                and val.strip().lower() in EMPTY_VALUES):
                            return val
        return None

    def normalize(self, raw_records: list[dict]) -> list[dict]:
        out = []
        ignored = 0
        for raw in raw_records:
            rec = {"source_city": self.city, "source_state": self.state,
                   "source_link": self.source_link, "city": self.city}
            clean = {k: v for k, v in raw.items() if k != "_dataset_label"}
            for std_field in self.STANDARD_FIELDS:
                rec[std_field] = self._pick(clean, std_field)
            if isinstance(rec.get("contractor"), str) and "," in rec["contractor"]:
                parts, seen = [], set()
                for part in (x.strip() for x in rec["contractor"].split(",")):
                    if part and part.lower() not in seen:
                        seen.add(part.lower())
                        parts.append(part)
                rec["contractor"] = ", ".join(parts)
            rec["estimated_value"] = self._to_float(rec["estimated_value"])
            # Valores como $0,01 ou $1 são marcadores administrativos, não o custo real da obra.
            if rec["estimated_value"] is not None and rec["estimated_value"] < 10:
                rec["estimated_value"] = None
            rec["issue_date"] = self._to_date(rec["issue_date"])
            if not rec.get("permit_type"):
                rec["permit_type"] = raw.get("_dataset_label")
            unknown_label = self.config.get("unknown_type_label")
            if unknown_label and (rec.get("permit_type") is None or
                                  str(rec["permit_type"]).strip().lower() in EMPTY_VALUES):
                rec["permit_type"] = unknown_label
            if rec.get("permit_number") is not None:
                rec["permit_number"] = str(rec["permit_number"])
            if not is_relevant(rec.get("permit_type"), rec.get("description")):
                ignored += 1
                continue
            rec["category"] = classify_permit(rec.get("permit_type"), rec.get("description"))
            out.append(rec)
        if ignored:
            print(f"[info] {self.city}: {ignored} permits de especialidade/administrativos ignorados "
                  f"(elétrico, hidráulico, gás, alarme, certificados...)")
        return [r for r in out if r.get("permit_number") and r.get("address")]

    def fetch(self, days_back: int = 2) -> list[dict]:
        return self.normalize(self.fetch_raw(days_back))

    @staticmethod
    def _to_float(v):
        if v is None:
            return None
        try:
            return float(str(v).replace(",", "").replace("$", ""))
        except (TypeError, ValueError):
            return None

    @staticmethod
    def _to_date(v):
        """Converte vários formatos de data para AAAA-MM-DD."""
        if v is None or v == "":
            return None
        if isinstance(v, datetime):
            return v.date().isoformat()
        if isinstance(v, date):
            return v.isoformat()
        if isinstance(v, (int, float)):
            seconds = v / 1000 if v > 1e11 else v
            try:
                return datetime.fromtimestamp(seconds, tz=timezone.utc).date().isoformat()
            except (OverflowError, OSError, ValueError):
                return None
        if isinstance(v, datetime):
            return v.date().isoformat()
        if isinstance(v, date):
            return v.isoformat()
        s = str(v).strip()
        if ISO_DATE.match(s):
            return s[:10]
        m = re.match(r"^(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})", s)
        if m:
            a, b, y = int(m.group(1)), int(m.group(2)), m.group(3)
            if len(y) == 2:
                y = "20" + y
            month, day = (b, a) if (a > 12 and b <= 12) else (a, b)   # padrão EUA: mês/dia/ano
            return f"{y}-{month:02d}-{day:02d}"
        return s


class CKANAdapter(BaseAdapter):
    """CKAN datastore_search_sql — usado por Boston (data.boston.gov)."""

    def _discover_date_field(self, base_url: str, resource_id: str) -> str:
        """Descobre a coluna da data de emissão (usado quando o cadastro diz date_field: auto)."""
        url = base_url.replace("datastore_search_sql", "datastore_search")
        data = _http_get(url, {"resource_id": resource_id, "limit": 1}).json()
        fields = [f["id"] for f in data.get("result", {}).get("fields", []) if f.get("id") != "_id"]
        lowered = {f.lower(): f for f in fields}
        for cand in DEFAULT_CANDIDATES["issue_date"] + ["issuedate", "issued", "date_issued", "issue_dt"]:
            if cand in lowered:
                print(f"[info] {self.city}: coluna de data escolhida: {lowered[cand]}")
                return lowered[cand]
        for f in fields:
            if "issue" in f.lower() and "date" in f.lower():
                print(f"[info] {self.city}: coluna de data escolhida: {f}")
                return f
        raise RuntimeError(f"CKAN: nenhuma coluna de data de emissão encontrada. Colunas: {fields}")

    def fetch_raw(self, days_back: int) -> list[dict]:
        base_url = self.config["base_url"]
        resource_id = self.config["resource_id"]
        date_field = self.config.get("date_field", "issued_date")
        if date_field == "auto":
            date_field = self._discover_date_field(base_url, resource_id)
        since = (date.today() - timedelta(days=days_back)).isoformat()
        sql = (f'SELECT * FROM "{resource_id}" WHERE {date_field} >= \'{since}\' '
               f'ORDER BY {date_field} DESC LIMIT 1000')
        data = _http_get(base_url, {"sql": sql}).json()
        if not data.get("success"):
            raise RuntimeError(f"CKAN API error ({self.city}): {data}")
        return data["result"]["records"]


class SocrataAdapter(BaseAdapter):
    """Socrata SODA API — usado por Cambridge e muitas outras cidades dos EUA."""

    def fetch(self, days_back: int = 2) -> list[dict]:
        records = self.normalize(self.fetch_raw(days_back))
        # Proteção final: só devolve permits cuja data de emissão seja recente
        # (evita tratar permits antigos como "novos" quando o filtro do servidor não pôde ser usado).
        since = (date.today() - timedelta(days=days_back)).isoformat()
        return [r for r in records if r.get("issue_date") and r["issue_date"] >= since]

    def _pick_date_field(self, columns: list[str], sample: dict) -> str | None:
        lowered = {c.lower(): c for c in columns}
        preferred = self.config.get("date_field")
        candidates = ([preferred] if preferred else []) + DEFAULT_CANDIDATES["issue_date"]
        chosen = None
        for name in candidates:
            if name and name.lower() in lowered:
                chosen = lowered[name.lower()]
                break
        if chosen is None:
            for c in columns:
                if "issue" in c.lower() and "date" in c.lower():
                    chosen = c
                    break
        # Só filtramos por data no servidor se a coluna for uma data de verdade (AAAA-MM-DD...).
        if chosen and ISO_DATE.match(str(sample.get(chosen, ""))):
            return chosen
        return None

    def _log_schema(self, label: str, dataset_id: str, rows: list[dict]):
        if not DEBUG:
            return
        if not rows:
            print(f"[info] {self.city} / {label} ({dataset_id}): sem linhas de amostra")
            return
        print(f"[info] {self.city} / {label} ({dataset_id}) COLUNAS: {sorted(rows[0].keys())}")
        sample = {k: (str(v)[:60]) for k, v in rows[0].items()}
        print(f"[info] {self.city} / {label} EXEMPLO: {sample}")

    def _log_freshness(self, url: str, label: str, date_field: str | None):
        """Escreve no log a data do permit mais recente do conjunto."""
        if not date_field:
            return
        try:
            res = _http_get(url, {"$select": f"max({date_field}) as latest"}).json()
            print(f"[info] {self.city} / {label}: permit mais recente emitido em: {res}")
        except Exception as e:  # noqa: BLE001
            print(f"[info] {self.city} / {label}: não consegui medir a data mais recente: {e}")

    def _log_types(self, url: str, label: str):
        """Escreve no log os tipos de permit mais frequentes do conjunto."""
        try:
            res = _http_get(url, {"$select": "permit_type, count(*) as n",
                                  "$group": "permit_type", "$order": "n DESC",
                                  "$limit": 40}).json()
            print(f"[info] {self.city} / {label}: TIPOS DE PERMIT: {res}")
        except Exception as e:  # noqa: BLE001
            print(f"[info] {self.city} / {label}: não consegui listar tipos: {e}")

    def fetch_raw(self, days_back: int) -> list[dict]:
        domain = self.config["domain"]
        since = (date.today() - timedelta(days=days_back)).isoformat()

        # Conjuntos só para "espiar" as colunas (não entram nos resultados). Só roda em modo depuração.
        inspect = (self.config.get("inspect_datasets") or {}) if DEBUG else {}
        for label, dataset_id in inspect.items():
            try:
                url = f"https://{domain}/resource/{dataset_id}.json"
                rows = _http_get(url, {"$limit": 2}).json()
                self._log_schema(label, dataset_id, rows)
                if rows:
                    self._log_freshness(url, label, self._pick_date_field(list(rows[0].keys()), rows[0]))
                self._log_types(url, label)
            except Exception as e:  # noqa: BLE001
                print(f"[info] não consegui espiar {label}: {e}")

        all_records: list[dict] = []
        for label, dataset_id in (self.config.get("datasets") or {}).items():
            url = f"https://{domain}/resource/{dataset_id}.json"
            try:
                sample_rows = _http_get(url, {"$limit": 2}).json()
                self._log_schema(label, dataset_id, sample_rows)
                if not sample_rows:
                    continue
                columns = list(sample_rows[0].keys())
                date_field = self._pick_date_field(columns, sample_rows[0])
                self._log_freshness(url, label, date_field)

                params = {"$limit": 1000}
                if date_field:
                    params["$where"] = f"{date_field} >= '{since}T00:00:00'"
                    params["$order"] = f"{date_field} DESC"
                else:
                    params["$order"] = ":updated_at DESC"
                    params["$limit"] = 200
                    print(f"[info] {self.city} / {label}: sem coluna de data utilizável; "
                          f"usando os 200 registros mais recentemente atualizados")

                try:
                    rows = _http_get(url, params).json()
                except RuntimeError as e:
                    print(f"[aviso] {self.city} / {label}: filtro falhou ({e}); tentando sem filtro")
                    rows = _http_get(url, {"$order": ":updated_at DESC", "$limit": 200}).json()

                print(f"[info] {self.city} / {label}: {len(rows)} registros recentes encontrados")
                for rec in rows:
                    rec["_dataset_label"] = label
                    all_records.append(rec)
            except Exception as e:  # noqa: BLE001
                print(f"[erro] {self.city} / {label} ({dataset_id}): {e}")
        return all_records


class ArcGISAdapter(BaseAdapter):
    """ArcGIS Hub (GeoJSON) — usado por Worcester e outras cidades."""

    def fetch_raw(self, days_back: int) -> list[dict]:
        geojson_url = self.config["geojson_url"]
        data = _http_get(geojson_url).json()
        features = data.get("features", [])
        rows = [f.get("properties", {}) for f in features]
        if rows:
            print(f"[info] {self.city} COLUNAS: {sorted(rows[0].keys())}")
            print(f"[info] {self.city} EXEMPLO: { {k: str(v)[:60] for k, v in rows[0].items()} }")
        return rows

    def fetch(self, days_back: int = 2) -> list[dict]:
        records = self.normalize(self.fetch_raw(days_back))
        since = (date.today() - timedelta(days=days_back)).isoformat()
        return [r for r in records if r.get("issue_date") and r["issue_date"] >= since]


class ArcGISQueryAdapter(BaseAdapter):
    """ArcGIS FeatureServer consultado pela API REST (/query), com paginação e espera
    automática quando o servidor limita as consultas (erro 429).

    Usado por Worcester. Config: service_url, order_by, page_size, max_pages, where.
    Ex.: em Worcester as datas são TEXTO (m/d/aaaa) e o registro mais novo vem primeiro
    (ObjectId crescente), então buscamos as primeiras páginas e filtramos a data aqui."""

    def _query(self, url: str, params: dict) -> dict:
        for attempt in range(6):
            resp = requests.get(url, params=params, headers=HEADERS, timeout=60)
            if resp.status_code == 429:
                time.sleep(30 * (attempt + 1))
                continue
            data = resp.json()
            err = data.get("error")
            if err and err.get("code") == 429:
                wait = 30 * (attempt + 1)
                print(f"[info] {self.city}: servidor pediu para esperar; aguardando {wait}s")
                time.sleep(wait)
                continue
            if err:
                raise RuntimeError(f"ArcGIS erro {err.get('code')}: {err.get('message')} {err.get('details')}")
            return data
        raise RuntimeError("ArcGIS: limite de consultas excedido após várias tentativas")

    def fetch_raw(self, days_back: int) -> list[dict]:
        base = self.config["service_url"].rstrip("/")
        page = int(self.config.get("page_size", 1000))
        pages = int(self.config.get("max_pages", 2))
        rows: list[dict] = []
        for i in range(pages):
            params = {"where": self.config.get("where", "1=1"), "outFields": "*",
                      "orderByFields": self.config.get("order_by", "ObjectId ASC"),
                      "resultOffset": i * page, "resultRecordCount": page, "f": "json"}
            feats = self._query(base + "/query", params).get("features", [])
            rows += [f.get("attributes", {}) for f in feats]
            if len(feats) < page:
                break
        return rows

    def fetch(self, days_back: int = 2) -> list[dict]:
        records = self.normalize(self.fetch_raw(days_back))
        since = (date.today() - timedelta(days=days_back)).isoformat()
        dated = [r["issue_date"] for r in records if r.get("issue_date")]
        if dated:
            print(f"[info] {self.city}: permit mais recente emitido em: {max(dated)}")
        return [r for r in records if r.get("issue_date") and r["issue_date"] >= since]


def _norm_header(h) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(h or "").lower()).strip("_")


class FileReportAdapter(BaseAdapter):
    """Cidades que publicam os permits emitidos em arquivos mensais (Excel ou CSV),
    como Reading. O adaptador abre a página de índice, acha os links dos arquivos mais
    recentes, baixa, localiza a linha de cabeçalho e traduz as colunas.

    Config: index_url, link_regex, link_text_regex, max_files, min_days_back.
    EXPERIMENTAL até o primeiro log confirmar os nomes reais das colunas."""

    HEADER_TOKENS = ("address", "permit", "date", "type", "contractor", "applicant", "cost")

    def _find_links(self, html: str) -> list[str]:
        link_re = re.compile(self.config.get("link_regex", r"DocumentCenter/View/\d+"), re.I)
        text_re = re.compile(self.config.get("link_text_regex", ""), re.I)
        picked: list[str] = []
        for href, inner in re.findall(r"<a[^>]+href=[\"']([^\"']+)[\"'][^>]*>(.*?)</a>", html, re.I | re.S):
            text = re.sub(r"<[^>]+>", " ", inner)
            if link_re.search(href) and text_re.search(text):
                url = urljoin(self.config["index_url"], href.replace("&amp;", "&"))
                if url not in picked:
                    picked.append(url)
        return picked[: int(self.config.get("max_files", 2))]

    def _parse_xlsx(self, content: bytes) -> list[dict]:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        out: list[dict] = []
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            hdr = None
            for i, r in enumerate(rows[:15]):
                cells = [str(c).strip() for c in r if c not in (None, "")]
                joined = " ".join(cells).lower()
                if len(cells) >= 3 and any(t in joined for t in self.HEADER_TOKENS):
                    hdr = i
                    break
            if hdr is None:
                continue
            headers = [_norm_header(c) for c in rows[hdr]]
            print(f"[info] {self.city}: colunas do arquivo ({ws.title}): {[h for h in headers if h]}")
            for r in rows[hdr + 1:]:
                if not any(c not in (None, "") for c in r):
                    continue
                out.append({h: v for h, v in zip(headers, r) if h})
        return out

    def _parse_csv(self, content: bytes) -> list[dict]:
        text = content.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        rows = [{_norm_header(k): v for k, v in row.items() if k} for row in reader]
        if rows:
            print(f"[info] {self.city}: colunas do arquivo: {list(rows[0].keys())}")
        return rows

    def fetch_raw(self, days_back: int) -> list[dict]:
        index_url = self.config["index_url"]
        if not robots_allows(index_url):
            raise RuntimeError(f"o robots.txt de {urlparse(index_url).netloc} não permite coleta automática desta página")
        html = _http_get(index_url).text
        urls = self._find_links(html)
        print(f"[info] {self.city}: {len(urls)} arquivo(s) de relatório encontrado(s)")
        rows: list[dict] = []
        for url in urls:
            try:
                if not robots_allows(url):
                    print(f"[aviso] {self.city}: robots.txt não permite baixar {url}; arquivo ignorado")
                    continue
                resp = _http_get(url)
                data = resp.content
                parsed = self._parse_xlsx(data) if data[:2] == b"PK" else self._parse_csv(data)
                print(f"[info] {self.city}: {len(parsed)} linhas em {url}")
                rows += parsed
            except Exception as e:  # noqa: BLE001
                print(f"[erro] {self.city}: falha ao ler {url}: {e}")
        return rows

    def fetch(self, days_back: int = 2) -> list[dict]:
        days = max(days_back, int(self.config.get("min_days_back", 45)))  # relatórios mensais têm atraso
        records = self.normalize(self.fetch_raw(days))
        since = (date.today() - timedelta(days=days)).isoformat()
        return [r for r in records if r.get("issue_date") and r["issue_date"] >= since]


class PermitEyesAdapter(BaseAdapter):
    """PermitEyes "Public View" (Full Circle Technologies) — usado por Taunton, Hingham, Falmouth,
    Concord, North Reading, Attleboro, Mansfield e cidades de Berkshire.

    A tela pública é uma tabela DataTables que pede os dados em ajax/<arquivo>.php. Este adaptador:
      1. abre a página pública e lê os nomes das colunas (que variam de cidade para cidade);
      2. acha o endereço de dados que responde em JSON;
      3. alinha as colunas com as células (alguns cabeçalhos não têm dado, ex.: Taunton);
      4. descobre se os registros mais novos ficam no começo ou no fim da lista;
      5. busca poucas páginas a partir do lado mais novo e traduz as colunas.
    Config: portal_url, ajax_urls (opcional), page_size, max_pages, newest_at ("start"/"end")."""

    DEFAULT_ENDPOINTS = ["ajax/getpublicview.php", "ajax/getbuildingpublichome.php"]
    ACTION_KEYS = {"application", "permit", "inspection", "app", "insp", "coc", "att", "co", "sign_off", "details"}
    _DATE = re.compile(r"^\d{1,2}[/-]\d{1,2}[/-]\d{2,4}")
    _MONEY = re.compile(r"^\$?[\d,]+(\.\d+)?$")
    _STREET = re.compile(r"\b(st|street|ave|avenue|rd|road|dr|drive|ln|lane|way|ct|court|pl|place|blvd|cir|circle|"
                         r"ter|terrace|hwy|pkwy|sq|square|path|trl|trail)\.?$", re.I)

    @staticmethod
    def _clean(cell) -> str:
        txt = re.sub(r"<[^>]+>", " ", str(cell if cell is not None else ""))
        return " ".join(html_lib.unescape(txt).split())

    def _tables(self, page: str) -> dict:
        out = {}
        for tid, body in re.findall(r"<table[^>]*\bid=[\"']([^\"']+)[\"'][^>]*>(.*?)</table>", page, re.I | re.S):
            heads = [self._clean(t) for t in re.findall(r"<th[^>]*>(.*?)</th>", body, re.I | re.S)]
            if heads:
                out[tid] = heads
        return out

    def _pick_headers(self, tables: dict, endpoint: str) -> list[str]:
        want = "building" if "building" in endpoint.lower() else "publicview"
        for tid, heads in tables.items():
            if want in tid.lower() and any("issue" in h.lower() for h in heads):
                return heads
        best = [h for h in tables.values() if any("issue" in x.lower() for x in h)]
        return max(best, key=len) if best else []

    def _session(self, view_url: str):
        sess = requests.Session()
        sess.headers.update(HEADERS)
        sess.headers.update({"X-Requested-With": "XMLHttpRequest", "Referer": view_url})
        return sess

    def _params(self, start: int, length: int, full: bool) -> dict:
        data = {"draw": 1, "start": start, "length": length, "search[value]": "", "search[regex]": "false"}
        data.update(self.config.get("extra_params") or {})
        if full:   # versões mais novas do sistema exigem a lista de colunas e a ordenação
            for i in range(self._ncols):
                data[f"columns[{i}][data]"] = i
                data[f"columns[{i}][name]"] = ""
                data[f"columns[{i}][searchable]"] = "true"
                data[f"columns[{i}][orderable]"] = "true"
                data[f"columns[{i}][search][value]"] = ""
                data[f"columns[{i}][search][regex]"] = "false"
            data["order[0][column]"] = self._order_col
            data["order[0][dir]"] = self._order_dir
        return data

    def _post(self, url: str, start: int, length: int, full: bool | None = None) -> dict:
        full = self._full if full is None else full
        last = None
        for attempt in range(3):
            try:
                resp = self.sess.post(url, data=self._params(start, length, full), timeout=60, allow_redirects=False)
                if 300 <= resp.status_code < 400:
                    raise RuntimeError(f"redirecionou para {resp.headers.get('Location', '?')[:60]} (exige sessão/login)")
                if resp.status_code >= 400:
                    raise RuntimeError(f"HTTP {resp.status_code} :: {' '.join(resp.text[:100].split())}")
                try:
                    js = resp.json()
                except ValueError:
                    visible = re.sub(r"<(script|style)[^>]*>.*?</\1>", "", resp.text, flags=re.S | re.I)
                    txt = " ".join(re.sub(r"<[^>]+>", " ", visible).split())[:160]
                    raise RuntimeError(f"resposta não é JSON :: {txt}")
                if not isinstance(js, dict) or "data" not in js:
                    raise RuntimeError("resposta sem o campo data")
                return js
            except RuntimeError as e:
                last = e
                if "redirecionou" in str(e) or "HTTP 4" in str(e):
                    break
                time.sleep(2 * (attempt + 1))
            except Exception as e:  # noqa: BLE001
                last = e
                time.sleep(2 * (attempt + 1))
        raise RuntimeError(f"{last}")

    def _resolve_town(self, page: str):
        """Páginas multi-cidade (Berkshire): cada aba tem data-town-id e data-url; o POST precisa do town_id."""
        tabs = []
        for m in re.finditer(r"<[^>]*data-town-id=[^>]*>", page, re.I):
            tag = m.group(0)
            inner = page[m.end(): m.end() + 80].split("<")[0]
            tid = re.search(r"data-town-id=[\"']?([\w-]+)", tag, re.I)
            url = re.search(r"data-url=[\"']([^\"']+)", tag, re.I)
            text = self._clean(inner)
            title = re.search(r"(?:title|data-town-name|aria-label)=[\"']([^\"']+)", tag + page[m.end(): m.end() + 200].split("</a>")[0], re.I)
            if tid:
                tabs.append({"id": tid.group(1), "url": url.group(1) if url else None,
                             "name": (title.group(1) if title else text)})
        if not tabs:
            return
        uniq = {(t["id"], t["name"]): t for t in tabs}.values()
        print(f"[info] {self.city}: abas de cidades encontradas: {[(t['id'], t['name'], t['url']) for t in list(uniq)[:14]]}")
        want_id = str(self.config.get("town_id") or "")
        want_name = str(self.config.get("town_name") or "").lower()
        chosen = None
        for t in uniq:
            if (want_id and t["id"] == want_id) or (want_name and want_name in t["name"].lower()):
                chosen = t
                break
        if chosen is None and len(set(t["id"] for t in uniq)) == 1:
            chosen = next(iter(uniq))
        if chosen:
            extra = dict(self.config.get("extra_params") or {})
            extra["town_id"] = chosen["id"]
            self.config = {**self.config, "extra_params": extra}
            if chosen["url"] and not self.config.get("ajax_urls"):
                self.config["ajax_urls"] = [chosen["url"]]
            print(f"[info] {self.city}: usando town_id={chosen['id']} ({chosen['name']})")
        else:
            print(f"[aviso] {self.city}: não achei a aba desta cidade; informe town_id ou town_name no cadastro")

    # ---- alinhamento das colunas ----
    def _plausible(self, key: str, cell: str) -> float:
        if not cell:
            return 0.2 if key in self.ACTION_KEYS else 0.0
        long_text = len(cell) >= 12 and not self._DATE.match(cell) and not self._MONEY.match(cell)
        if key in self.ACTION_KEYS:
            return -1.0 if long_text else -0.5
        if "description" in key or key.startswith("brief"):
            return 0.8 if long_text else 0.0
        if key.startswith("col") and key[3:].isdigit():       # cabeçalho sem nome
            return -0.5
        if "date" in key:
            return 1.0 if self._DATE.match(cell) else -1.5
        if "cost" in key or "valu" in key:
            return 1.0 if self._MONEY.match(cell) else -1.5
        if key == "ap_no":
            return 1.0 if cell.isdigit() else -1.0
        if "permit_number" in key:
            return 1.0 if re.match(r"^[A-Za-z]{0,6}[-\s]?\d", cell) else -1.0
        if key == "street_name":
            return 1.0 if self._STREET.search(cell) else 0.0
        if key == "appl_type":
            if re.search(r"\d{3,}", cell):
                return -1.0
            return 0.6 if re.fullmatch(r"[A-Za-z .()/&-]{2,16}", cell) else -0.3
        if "status" in key:
            return 1.0 if re.fullmatch(r"[A-Za-z .-]{3,20}", cell) else -1.0
        if key == "site_address":
            return 1.0 if re.match(r"^\d+\s*\w", cell) else -0.5
        if key == "street_no":
            return 0.5 if re.match(r"^\d+\w?$", cell) else -0.5
        if key in ("applicant", "owner", "contractor_name"):
            if self._DATE.match(cell) or self._MONEY.match(cell):
                return -1.0
            if self._STREET.search(cell):                       # nome de rua num campo de pessoa
                return -0.8
            if re.fullmatch(r"[A-Z][A-Z .()/-]{1,7}", cell):    # sigla de tipo (RESI, ELECT...) num campo de pessoa
                return -0.6
            return 0.0
        return 0.0

    def _build_alignment(self, keys: list[str], sample_rows: list[list[str]]) -> list[str | None]:
        """Devolve, para cada célula, o nome da coluna (ou None). Testa quais cabeçalhos não têm dado."""
        from itertools import combinations
        sample_rows = [r for r in sample_rows if r]
        if not sample_rows:
            return list(keys)
        counts: dict[int, int] = {}
        for r in sample_rows:
            counts[len(r)] = counts.get(len(r), 0) + 1
        n_cells = max(counts, key=counts.get)
        H = len(keys)
        if n_cells == H:
            return list(keys)
        if n_cells > H:
            return [None] * (n_cells - H) + list(keys)
        d = H - n_cells
        forced = [str(x) for x in (self.config.get("drop_headers") or [])]
        if forced:    # a cidade informa explicitamente quais cabeçalhos não trazem dado
            kept = [k for k in keys if k not in forced]
            if len(kept) == n_cells:
                print(f"[info] {self.city}: cabeçalhos sem dado (definidos no cadastro): {forced}")
                return kept
        if d > 3:
            return list(keys[H - n_cells:])
        rows = [r for r in sample_rows if len(r) == n_cells][:40]
        best, best_score = None, None
        for omit in combinations(range(H), d):
            kept = [k for i, k in enumerate(keys) if i not in omit]
            score = sum(self._plausible(k, c) for r in rows for k, c in zip(kept, r))
            if best_score is None or score > best_score:
                best, best_score = kept, score
        print(f"[info] {self.city}: {H} cabeçalhos para {n_cells} células; cabeçalhos sem dado: "
              f"{[k for k in keys if k not in best]}")
        return best

    def _to_dicts(self, rows: list, keys: list[str], aligned: list) -> list[dict]:
        out = []
        for r in rows:
            cells = [self._clean(c) for c in r]
            if len(cells) == len(aligned):
                pairs = zip(aligned, cells)
            else:
                n = min(len(cells), len(keys))
                pairs = zip(keys[len(keys) - n:], cells[len(cells) - n:])
            d = {k: v for k, v in pairs if k and v}
            if "site_address" not in d and (d.get("street_no") or d.get("street_name")):
                d["site_address"] = f"{d.get('street_no', '')} {d.get('street_name', '')}".strip()
            out.append(d)
        return out

    def _row_date(self, d: dict) -> str:
        v = self._to_date(self._pick(d, "issue_date"))
        return v if isinstance(v, str) and ISO_DATE.match(v) else ""

    def fetch_raw(self, days_back: int) -> list[dict]:
        view_url = self.config["portal_url"]
        if not robots_allows(view_url):
            raise RuntimeError("o robots.txt não permite coleta automática desta página")
        self.sess = self._session(view_url)
        resp = self.sess.get(view_url, timeout=60)
        if resp.status_code >= 400:
            raise RuntimeError(f"HTTP {resp.status_code} ao abrir a página pública")
        page = resp.text
        tables = self._tables(page)
        self._resolve_town(page)

        endpoints = list(self.config.get("ajax_urls") or self.DEFAULT_ENDPOINTS)
        for ep in re.findall(r"ajax/[A-Za-z_]+\.php", page):
            if ep not in endpoints and not re.search(r"attach|modal|inspect|check", ep, re.I):
                endpoints.append(ep)

        self._ncols = max([len(h) for h in tables.values()] or [20])
        self._order_col = 0
        self._order_dir = "asc"
        self._full = False
        url = js0 = None
        tried = []
        for ep in endpoints:
            full_url = urljoin(view_url, ep)
            for mode in (False, True):      # primeiro o pedido simples; depois o completo do DataTables
                try:
                    js0 = self._post(full_url, 0, 25, full=mode)
                except Exception as e:  # noqa: BLE001
                    tried.append(f"{ep} ({'completo' if mode else 'simples'}) -> {e}")
                    continue
                if js0.get("data") or js0.get("recordsTotal"):
                    url, self._full = full_url, mode
                    break
                tried.append(f"{ep} ({'completo' if mode else 'simples'}) -> vazio")
            if url:
                break
        if not url:
            for t in tried[:10]:
                print(f"[info] {self.city}: tentativa {t}")
            raise RuntimeError(f"nenhum endereço de dados respondeu ({len(tried)} tentativas; veja as mensagens)")

        heads = self._pick_headers(tables, url)
        print(f"[info] {self.city}: endereço de dados {url.split('/')[-1]} (pedido {'completo' if self._full else 'simples'}); colunas: {heads}")
        if not heads:
            raise RuntimeError("não encontrei os nomes das colunas na página")
        keys = [_norm_header(h) or f"col{i}" for i, h in enumerate(heads)]

        total = int(str(js0.get("recordsTotal") or js0.get("recordsFiltered") or len(js0.get("data", []))).replace(",", "") or 0)
        sample = [[self._clean(c) for c in r] for r in js0.get("data", [])]
        last_js = None
        newest_at = self.config.get("newest_at")
        if total > 25:
            last_js = self._post(url, max(total - 25, 0), 25)
            sample += [[self._clean(c) for c in r] for r in last_js.get("data", [])]
        aligned = self._build_alignment(keys, sample)

        ordered = False
        if self._full and not newest_at and "issue_date" in aligned:
            self._order_col, self._order_dir = aligned.index("issue_date"), "desc"
            try:
                chk = self._to_dicts(self._post(url, 0, 25).get("data", []), keys, aligned)
                vals = [d for d in (self._row_date(r) for r in chk) if d]
                if len(vals) >= 3 and vals == sorted(vals, reverse=True):
                    newest_at, ordered = "start", True
                    print(f"[info] {self.city}: ordenação por Issue Date (decrescente) confirmada; mais novos no início")
            except Exception as e:  # noqa: BLE001
                print(f"[info] {self.city}: ordenação por data não funcionou ({e})")
            if not ordered:
                self._order_col, self._order_dir = 0, "asc"

        first = self._to_dicts(js0.get("data", []), keys, aligned)
        if not newest_at:
            last = self._to_dicts(last_js.get("data", []), keys, aligned) if last_js else first
            d_first = max([d for d in (self._row_date(r) for r in first) if d] or [""])
            d_last = max([d for d in (self._row_date(r) for r in last) if d] or [""])
            if d_first and d_last and d_first != d_last:
                newest_at = "end" if d_last > d_first else "start"
            else:
                newest_at = "end"          # em todas as cidades verificadas, os mais novos ficam no fim
            print(f"[info] {self.city}: {total} registros; data no início={d_first or '?'}, no fim={d_last or '?'}; mais novos no {newest_at}")
        if first:
            print(f"[info] {self.city}: exemplo de linha: {first[0]}")

        since = (date.today() - timedelta(days=days_back)).isoformat()
        size = int(self.config.get("page_size", 100))
        rows: list[dict] = []
        pos_end, pos_start = total, 0
        for _ in range(int(self.config.get("max_pages", 6))):
            if newest_at == "start":
                st, length = pos_start, min(size, max(total - pos_start, 1))
            else:
                st = max(pos_end - size, 0)
                length = pos_end - st
            if length <= 0:
                break
            js = self._post(url, st, length)
            got = js.get("data", [])
            if 0 < len(got) < length and size > len(got):   # o servidor limita o tamanho da página
                size = len(got)
                if newest_at == "end":
                    continue
            batch = self._to_dicts(got, keys, aligned)
            rows += batch
            if not got:
                break
            dates = [d for d in (self._row_date(b) for b in batch) if d]
            if dates and min(dates) < since:
                break
            if newest_at == "start":
                pos_start += len(got)
                if pos_start >= total:
                    break
            else:
                pos_end = st
                if pos_end <= 0:
                    break
            time.sleep(1)

        # Diagnóstico: quais tipos (siglas) a cidade usa nos registros lidos
        tcount: dict[str, int] = {}
        for r in rows:
            t = self._pick(r, "permit_type")
            if t:
                tcount[str(t)] = tcount.get(str(t), 0) + 1
        top = sorted(tcount.items(), key=lambda kv: -kv[1])[:14]
        print(f"[info] {self.city}: tipos nas {len(rows)} linhas lidas: {top}")
        return rows

    def fetch(self, days_back: int = 2) -> list[dict]:
        records = self.normalize(self.fetch_raw(days_back))
        since = (date.today() - timedelta(days=days_back)).isoformat()
        dated = [r["issue_date"] for r in records if r.get("issue_date")]
        if dated:
            print(f"[info] {self.city}: data mais recente encontrada: {max(dated)}")
        return [r for r in records if r.get("issue_date") and r["issue_date"] >= since]


ADAPTERS = {
    "ckan": CKANAdapter,
    "socrata": SocrataAdapter,
    "arcgis": ArcGISAdapter,
    "arcgis_query": ArcGISQueryAdapter,
    "file_report": FileReportAdapter,
    "permiteyes": PermitEyesAdapter,
}
